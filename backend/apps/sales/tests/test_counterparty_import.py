import io
from datetime import timedelta
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.core.models import AuditLog
from apps.permissions.models import Permission
from apps.sales.models import Counterparty, CounterpartyImportJob, DebtPayment, Sale
from apps.sales.services.counterparty_service import CounterpartyService, CounterpartyServiceError
from tests.factories import CounterpartyFactory, EmployeeFactory, StaffFactory, TenantFactory

pytestmark = pytest.mark.django_db
ROOT = "/api/v1/sales/counterparties/import/"


@pytest.fixture
def client(owner, tenant):
    client = APIClient()
    client.force_authenticate(owner)
    return client


def xlsx(rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return SimpleUploadedFile("mijozlar.xlsx", output.getvalue())


def csv(body):
    return SimpleUploadedFile("mijozlar.csv", body.encode("utf-8-sig"))


def commit_url(preview):
    assert preview.status_code == 201, preview.data
    return ROOT + f"{preview.data['id']}/commit/"


def test_xlsx_preview_commit_balances_history_and_retry(client, owner, tenant):
    preview = client.post(
        ROOT + "preview/",
        {
            "file": xlsx(
                [
                    ["mijoz", "telefon", "nasiya_uzs", "nasiya_usd", "izoh"],
                    ["Ali Valiyev", "+998 (90) 123-45-67", "1 250 000,50", "75,25", "Eski daftar"],
                    ["Nodira", 998901234568, None, None, "Oddiy mijoz"],
                ]
            )
        },
        format="multipart",
    )
    url = commit_url(preview)
    assert not any(row["errors"] for row in preview.data["rows"])
    assert not Counterparty.objects.filter(tenant=tenant).exists()
    assert not DebtPayment.objects.filter(tenant=tenant).exists()
    from apps.cashbox.models import CashShift

    # Opening debt does not require an open shift or an exchange rate.
    CashShift.objects.filter(tenant=tenant).update(closed_at=timezone.now())
    result = client.post(url)
    assert result.status_code == 200, result.data
    assert result.data["created_count"] == 2
    assert Decimal(result.data["total_debt_uzs"]) == Decimal("1250000.50")
    assert Decimal(result.data["total_debt_usd"]) == Decimal("75.25")
    cp = Counterparty.objects.get(tenant=tenant, name="Ali Valiyev")
    assert cp.phone_number == "+998901234567"
    assert cp.debt_balance_uzs == Decimal("1250000.50")
    assert cp.debt_balance_usd == Decimal("75.25")
    assert cp.note == "Eski daftar"
    assert set(cp.payments.values_list("amount", "currency")) == {
        (Decimal("-1250000.50"), "UZS"),
        (Decimal("-75.25"), "USD"),
    }
    for payment in cp.payments.all():
        assert payment.is_correction and payment.payment_type == "none"
        assert payment.shift_id is None and payment.exchange_rate is None and payment.converted_amount_uzs is None
        assert payment.recorded_by == owner and payment.branch_id == tenant.id
    assert not Sale.objects.filter(tenant=tenant).exists()
    assert client.post(url).data == result.data
    assert Counterparty.objects.filter(tenant=tenant).count() == 2
    assert DebtPayment.objects.filter(tenant=tenant).count() == 2
    assert AuditLog.objects.filter(tenant=tenant, action="counterparty_import").count() == 1
    listing = client.get("/api/v1/sales/counterparties/?page=1")
    assert listing.data["summary"]["debtors_count"] == 1
    assert listing.data["summary"]["settled_count"] == 1
    assert Decimal(listing.data["summary"]["total_debt_usd"]) == Decimal("75.25")
    history = client.get(f"/api/v1/sales/counterparties/{cp.pk}/payments/")
    assert len(history.data) == 2 and all(row["is_correction"] for row in history.data)


def test_downloaded_template_can_be_filled_and_imported(client):
    response = client.get(ROOT + "template/")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Mijozlar", "Qo‘llanma"]
    assert list(next(workbook.active.values)) == ["mijoz", "telefon", "nasiya_uzs", "nasiya_usd", "izoh"]
    assert workbook.active.column_dimensions["B"].number_format == "@"
    workbook.active.append(["Shablondan", "+998901000001", 50000, 20.15, "Namuna"])
    output = io.BytesIO()
    workbook.save(output)
    preview = client.post(ROOT + "preview/", {"file": SimpleUploadedFile("shablon.xlsx", output.getvalue())})
    assert client.post(commit_url(preview)).status_code == 200


def test_csv_defaults_and_other_tenant_phone_are_allowed(client, tenant):
    CounterpartyFactory(tenant=TenantFactory(), phone_number="+998901234567")
    preview = client.post(ROOT + "preview/", {"file": csv("mijoz;telefon\nAli;998901234567\n")})
    assert client.post(commit_url(preview)).status_code == 200
    cp = Counterparty.objects.get(tenant=tenant)
    assert cp.debt_balance_usd == cp.debt_balance_uzs == 0
    assert not cp.payments.exists()


def test_imported_debt_can_be_paid_and_opening_balance_is_not_revenue(client, tenant):
    from apps.analytics.services.analytics_service import AnalyticsService

    preview = client.post(
        ROOT + "preview/", {"file": csv("mijoz,telefon,nasiya_uzs,nasiya_usd\nAli,+998901000001,100000,50\n")}
    )
    assert client.post(commit_url(preview)).status_code == 200
    cp = Counterparty.objects.get(tenant=tenant)
    summary = AnalyticsService.get_dashboard_summary(tenant=tenant)
    assert summary["kpi"]["total_revenue_uzs"] == "0.00"
    assert summary["kpi"]["total_revenue_usd"] == "0.00"
    response = client.post(
        f"/api/v1/sales/counterparties/{cp.pk}/payments/",
        {"amount": "10.00", "currency": "USD", "payment_type": "cash"},
        format="json",
    )
    assert response.status_code == 201, response.data
    cp.refresh_from_db()
    assert cp.debt_balance_uzs == Decimal("100000.00")
    assert cp.debt_balance_usd == Decimal("40.00")
    assert cp.payments.filter(is_correction=False).count() == 1
    summary = AnalyticsService.get_dashboard_summary(tenant=tenant)
    assert summary["kpi"]["total_revenue_uzs"] == "0.00"
    assert summary["kpi"]["total_revenue_usd"] == "10.00"
    # The timeline exposes the import and its accounting entries under finance.
    history = client.get("/api/v1/audit/history/?category=finance")
    assert history.status_code == 200, history.data
    assert any(row.get("action") == "counterparty_import" for row in history.data["results"])


def test_duplicates_archived_records_and_invalid_amounts_block_whole_file(client, tenant):
    CounterpartyFactory(tenant=tenant, phone_number="+998 (90) 100-00-01", is_active=False)
    preview = client.post(
        ROOT + "preview/",
        {
            "file": xlsx(
                [
                    ["mijoz", "telefon", "nasiya_uzs", "nasiya_usd"],
                    ["Valid", "+998901000002", 500, 0],
                    ["Takror", "998901000002", 0, 0],
                    ["Arxiv", "+998901000001", 0, 0],
                    ["Manfiy", "+998901000003", -10, 0],
                    ["Kasr", "+998901000004", 0, "1.001"],
                    ["Telefon", "abc", 0, 0],
                    ["Katta", "+998901000005", "1000000000000", 0],
                ]
            )
        },
    )
    url = commit_url(preview)
    assert not preview.data["rows"][0]["errors"]
    for row in preview.data["rows"][1:]:
        assert row["errors"]
    assert client.post(url).status_code == 400
    assert Counterparty.objects.filter(tenant=tenant).count() == 1
    assert not DebtPayment.objects.filter(tenant=tenant).exists()


def test_commit_revalidates_duplicates_and_expires(client, tenant):
    preview = client.post(ROOT + "preview/", {"file": csv("mijoz,telefon\nAli,+998901000002\n")})
    url = commit_url(preview)
    CounterpartyFactory(tenant=tenant, phone_number="+998901000002")
    response = client.post(url)
    assert response.status_code == 400
    assert "phone_number" in response.data["rows"][0]["errors"]
    CounterpartyImportJob.objects.filter(pk=preview.data["id"]).update(expires_at=timezone.now() - timedelta(seconds=1))
    assert client.post(url).status_code == 400


def test_staff_permissions_and_job_scope(client, tenant):
    preview = client.post(ROOT + "preview/", {"file": csv("mijoz,telefon\nAli,+998901000002\n")})
    url = commit_url(preview)
    staff = StaffFactory()
    employee = EmployeeFactory(user=staff, tenant=tenant)
    client.force_authenticate(staff)
    assert client.get(ROOT + "template/").status_code == 403
    assert client.post(ROOT + "preview/", {"file": csv("mijoz,telefon\nA,+998901000002\n")}).status_code == 403
    employee.permissions.add(Permission.objects.get(category="sales", codename="manage_counterparty"))
    assert client.post(url).status_code == 404
    assert client.get(ROOT + "template/").status_code == 200
    preview = client.post(ROOT + "preview/", {"file": csv("mijoz,telefon,nasiya_uzs\nA,+998901000002,10\n")})
    assert "debt_balance_uzs" in preview.data["rows"][0]["errors"]
    assert client.post(commit_url(preview)).status_code == 400
    permission = Permission.objects.get(category="sales", codename="record_debt_payment")
    employee.permissions.add(permission)
    preview = client.post(ROOT + "preview/", {"file": csv("mijoz,telefon,nasiya_usd\nA,+998901000002,10\n")})
    assert preview.data["rows"][0]["errors"] == {}
    employee.permissions.remove(permission)
    assert client.post(commit_url(preview)).status_code == 400
    preview = client.post(ROOT + "preview/", {"file": csv("mijoz,telefon\nA,+998901000002\n")})
    employee.permissions.clear()
    assert client.post(commit_url(preview)).status_code == 403
    other = TenantFactory()
    assert client.post(url, HTTP_X_TENANT_ID=str(other.pk)).status_code == 403
    EmployeeFactory(user=other.owner, tenant=other, is_owner=True)
    client.force_authenticate(other.owner)
    assert client.post(url, HTTP_X_TENANT_ID=str(other.pk)).status_code == 404
    assert not Counterparty.objects.filter(tenant=tenant).exists()


@pytest.mark.parametrize(
    "rows",
    [
        [["mijoz", "mijoz", "telefon"], ["A", "B", "+998901000001"]],
        [["mijoz", "telefon", "boshqa"], ["A", "+998901000001", 1]],
        [["mijoz", "nasiya_uzs"], ["A", 10]],
        [["mijoz", "telefon", "nasiya_uzs"], ["A", "+998901000001", "=100+50"]],
        [["mijoz", "telefon"]],
        [["mijoz", "telefon"], *[["A", "+998901000001"] for _ in range(1001)]],
    ],
)
def test_invalid_files_are_rejected(client, rows):
    assert client.post(ROOT + "preview/", {"file": xlsx(rows)}).status_code == 400


def test_file_limits_and_invalid_archives(client):
    for upload in [
        SimpleUploadedFile("bad.xlsx", b"not an archive"),
        SimpleUploadedFile("old.xls", b"xls"),
        SimpleUploadedFile("big.xlsx", b"a" * (8 * 1024 * 1024 + 1)),
    ]:
        assert client.post(ROOT + "preview/", {"file": upload}).status_code == 400


def test_mid_import_failure_rolls_back_everything(client, tenant, monkeypatch):
    preview = client.post(
        ROOT + "preview/", {"file": csv("mijoz,telefon,nasiya_uzs\nAli,+998901000001,100\nVali,+998901000002,200\n")}
    )
    create = CounterpartyService.create

    def fail_second(**kwargs):
        if kwargs["name"] == "Vali":
            raise CounterpartyServiceError("Sinov xatosi")
        return create(**kwargs)

    monkeypatch.setattr(CounterpartyService, "create", fail_second)
    assert client.post(commit_url(preview)).status_code == 400
    assert not Counterparty.objects.filter(tenant=tenant).exists()
    assert not DebtPayment.objects.filter(tenant=tenant).exists()
    assert not AuditLog.objects.filter(tenant=tenant, action="counterparty_import").exists()
    assert CounterpartyImportJob.objects.get(pk=preview.data["id"]).completed_at is None
