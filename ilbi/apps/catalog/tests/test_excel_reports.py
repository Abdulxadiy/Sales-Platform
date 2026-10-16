import io
from datetime import datetime
from datetime import timezone as dt_timezone
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.catalog.services.import_service import TEMPLATE_HEADERS
from apps.core.spreadsheets import INTEGER_FORMAT, MONEY_FORMAT, QUANTITY_FORMAT, XLSX_MIME, excel_response
from apps.inventory.models import Stock
from apps.permissions.models import Permission
from apps.sales.models import Sale
from tests.factories import (
    CategoryFactory,
    CounterpartyFactory,
    EmployeeFactory,
    PlatformAdminFactory,
    ProductFactory,
    ProductVariantFactory,
    SaleFactory,
    StaffFactory,
    TenantFactory,
)

pytestmark = pytest.mark.django_db


@pytest.fixture
def client(owner, tenant):
    client = APIClient()
    client.force_authenticate(owner)
    return client


def workbook(response):
    assert response.status_code == 200, getattr(response, "data", "")
    assert response["Content-Type"] == XLSX_MIME
    assert response["Cache-Control"] == "private, no-store"
    return load_workbook(io.BytesIO(response.content))


@pytest.fixture
def inventory(tenant):
    records = []
    for currency, unit, quantity, price, cost in [
        ("UZS", "dona", "3", "1000.25", "700.50"),
        ("USD", "kg", "1.125", "12.40", "8.80"),
        ("UZS", "dona", "2", "500", None),
    ]:
        category = CategoryFactory(tenant=tenant, currency=currency)
        product = ProductFactory(tenant=tenant, category=category)
        variant = ProductVariantFactory(
            tenant=tenant,
            product=product,
            unit=unit,
            barcode=f"001234{len(records)}",
            price_recommended=Decimal(price),
        )
        records.append(
            Stock.objects.create(
                tenant=tenant,
                product_variant=variant,
                quantity=Decimal(quantity),
                last_cost_price=Decimal(cost) if cost else None,
            )
        )
    return records


def test_stock_report_formats_identifiers_numbers_and_print_layout(client, inventory, tenant):
    tenant.name = '=HYPERLINK("untrusted")'
    tenant.save(update_fields=["name"])
    book = workbook(client.get("/api/v1/inventory/stock/export/", {"scope": "all"}))
    sheet = book["Ombor"]
    assert book.active == sheet
    assert sheet.freeze_panes == "B2"
    assert sheet.auto_filter.ref == f"A1:P{len(inventory) + 1}"
    assert not sheet.sheet_view.showGridLines
    assert sheet["A1"].font.bold and sheet["A1"].font.color.rgb == "00FFFFFF"
    assert sheet["A1"].fill.fgColor.rgb == "0017324D"
    assert sheet["A3"].fill.fgColor.rgb == "00F0F5FA"
    assert sheet["D2"].data_type == "s" and sheet["D2"].number_format == "@"
    assert sheet["E2"].value.startswith("00") and sheet["E2"].number_format == "@"
    assert sheet["G2"].data_type == "n" and sheet["G2"].number_format == INTEGER_FORMAT
    assert sheet["G1"].alignment.horizontal == sheet["G2"].alignment.horizontal == "center"
    assert sheet["I1"].alignment.horizontal == sheet["I2"].alignment.horizontal == "right"
    assert sheet["A1"].alignment.horizontal == sheet["A2"].alignment.horizontal == "left"
    assert sheet["G2"].border.left.style == sheet["G2"].border.right.style == "thin"
    assert sheet["I2"].number_format == MONEY_FORMAT
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 0
    assert sheet.print_title_rows == "$1:$1"
    summary = dict(book["Hisobot"].values)
    assert summary["Qatorlar soni"] == 3
    assert summary["Qoldiq · dona"] == 5
    assert summary["Qoldiq · kg"] == 1.125
    assert summary["Sotuv qiymati · UZS"] == 4000.75
    assert summary["Sotuv qiymati · USD"] == 13.95
    assert summary["Tannarx qiymati · UZS"] == 2101.5
    assert summary["Tannarx qiymati · USD"] == 9.9
    assert book["Hisobot"]["B2"].data_type == "s"
    assert book["Hisobot"]["B2"].value == tenant.name


@pytest.mark.parametrize("allow_cost", [False, True])
def test_staff_cannot_export_even_with_stock_and_cost_permissions(client, inventory, tenant, allow_cost):
    staff = StaffFactory()
    employee = EmployeeFactory(user=staff, tenant=tenant)
    employee.permissions.add(Permission.objects.get(category="inventory", codename="view_stock"))
    if allow_cost:
        employee.permissions.add(Permission.objects.get(category="catalog", codename="view_cost_price"))
    client.force_authenticate(staff)
    assert client.get("/api/v1/inventory/stock/export/", {"scope": "all"}).status_code == 403
    employee.permissions.clear()
    assert client.get("/api/v1/inventory/stock/export/", {"scope": "all"}).status_code == 403


@pytest.mark.parametrize("role", ["staff", "platform_admin", "other_owner"])
def test_all_excel_downloads_reject_nonowners_including_direct_requests(client, tenant, role):
    if role == "staff":
        user = StaffFactory()
        employee = EmployeeFactory(user=user, tenant=tenant)
        employee.permissions.set(Permission.objects.all())
    elif role == "platform_admin":
        user = PlatformAdminFactory()
    else:
        other = TenantFactory()
        user = other.owner
        EmployeeFactory(user=user, tenant=other)
    client.force_authenticate(user)
    for endpoint in [
        "/api/v1/inventory/stock/export/",
        "/api/v1/inventory/stock/export/?scope=all",
        "/api/v1/sales/export/",
        "/api/v1/sales/counterparties/export/",
        "/api/v1/catalog/import/template/",
    ]:
        assert client.get(endpoint, HTTP_X_TENANT_ID=str(tenant.pk)).status_code == 403, endpoint
        assert client.head(endpoint, HTTP_X_TENANT_ID=str(tenant.pk)).status_code == 403, endpoint
        assert client.options(endpoint, HTTP_X_TENANT_ID=str(tenant.pk)).status_code == 403, endpoint
    for endpoint in ["/api/v1/inventory/stock/export/", "/api/v1/sales/counterparties/export/"]:
        response = client.post(endpoint, {"ids": []}, format="json", HTTP_X_TENANT_ID=str(tenant.pk))
        assert response.status_code == 403, endpoint


def test_long_product_names_wrap_and_expand_their_own_rows(client, inventory):
    product = inventory[0].product_variant.product
    product.name = "11D Full Glue Himoya oynasi Redmi 12 Matoviy " * 4 + "(Model #03478)"
    product.save(update_fields=["name"])
    book = workbook(client.get("/api/v1/inventory/stock/export/", {"scope": "all"}))
    sheet = book.active
    row = next(row for row in sheet.iter_rows(min_row=2) if row[0].value == product.name)
    assert row[0].value.endswith("(Model #03478)")
    assert row[0].alignment.wrap_text
    assert row[0].font.bold
    assert sheet.column_dimensions["A"].width >= 60
    assert sheet.row_dimensions[row[0].row].height >= 78
    assert sheet.row_dimensions[row[0].row].height > sheet.row_dimensions[row[0].row + 1].height


def test_quantity_cells_show_zero_integers_and_fractions_without_dangling_separator():
    quantities = [Decimal("0.000"), Decimal("181.000"), Decimal("1.125"), Decimal("2.500")]
    book = workbook(
        excel_response(
            ["Mahsulot", "Qoldiq"],
            ([f"Item {index}", quantity] for index, quantity in enumerate(quantities)),
            "test.xlsx",
            formats={"Qoldiq": QUANTITY_FORMAT},
        )
    )
    cells = [row[1] for row in book.active.iter_rows(min_row=2)]
    assert [cell.value for cell in cells] == [0, 181, 1.125, 2.5]
    assert all(cell.data_type == "n" for cell in cells)
    assert [cell.number_format for cell in cells] == [
        INTEGER_FORMAT,
        INTEGER_FORMAT,
        QUANTITY_FORMAT,
        "#,##0.0;[Red](#,##0.0);0",
    ]
    assert "—" not in MONEY_FORMAT
    assert all(cell.alignment.horizontal == "center" for cell in cells)


def test_debt_report_keeps_negative_balances_and_currency_totals(client, tenant):
    CounterpartyFactory(
        tenant=tenant, debt_balance_uzs="100.50", debt_balance_usd="2.25", phone_number="+998901234567", name="Client A"
    )
    CounterpartyFactory(tenant=tenant, debt_balance_uzs="-20.25", debt_balance_usd="-0.50")
    CounterpartyFactory(debt_balance_uzs="99999")
    book = workbook(client.get("/api/v1/sales/counterparties/export/"))
    assert len(list(book.active.values)) == 3
    summary = dict(book["Hisobot"].values)
    assert summary["Qarz UZS · Jami"] == 80.25
    assert summary["Qarz USD · Jami"] == 1.75
    phone = next(row[1] for row in book.active.iter_rows(min_row=2) if row[0].value == "Client A")
    assert phone.value == "+998901234567"
    assert phone.data_type == "s"


def test_sales_report_totals_by_currency_and_preserves_tashkent_time(client, tenant, owner):
    first = SaleFactory(tenant=tenant, sold_by=owner, currency="USD", total_amount="10.25", refunded_amount="2.50")
    Sale.objects.filter(pk=first.pk).update(created_at=datetime(2026, 1, 1, 22, 30, tzinfo=dt_timezone.utc))
    SaleFactory(tenant=tenant, sold_by=owner, currency="UZS", total_amount="25000", refunded_amount="5000")
    SaleFactory(total_amount="99999")
    book = workbook(client.get("/api/v1/sales/export/"))
    summary = dict(book["Hisobot"].values)
    assert summary["Qatorlar soni"] == 2
    assert summary["Asl summa · USD"] == 10.25
    assert summary["Qaytarilgan · USD"] == 2.5
    assert summary["Sof summa · USD"] == 7.75
    assert summary["Sof summa · UZS"] == 20000
    data = list(book.active.values)
    row = next(row for row in data[1:] if row[0] == first.receipt_number)
    assert row[1] == datetime(2026, 1, 2, 3, 30)
    filtered = workbook(client.get("/api/v1/sales/export/", {"currency": "USD"}))
    summary = dict(filtered["Hisobot"].values)
    assert summary["Qatorlar soni"] == 1 and summary["Valyuta"] == "USD"
    assert "Sof summa · UZS" not in summary


def test_empty_report_and_import_template_remain_usable(client):
    book = workbook(client.get("/api/v1/inventory/stock/export/", {"scope": "all"}))
    assert len(list(book.active.values)) == 1
    assert dict(book["Hisobot"].values)["Qatorlar soni"] == 0
    template = workbook(client.get("/api/v1/catalog/import/template/"))
    assert list(next(template.active.values)) == TEMPLATE_HEADERS
    assert template.active.max_row == 1
    assert "Qo‘llanma" in template.sheetnames
    template.active.append(["Example", "Test", "Standart", "kg", "0012345678905", 0, 0, 100, "UZS", 0, 0])
    output = io.BytesIO()
    template.save(output)
    preview = client.post(
        "/api/v1/catalog/import/preview/",
        {
            "file": SimpleUploadedFile("template.xlsx", output.getvalue(), content_type=XLSX_MIME),
        },
    )
    assert preview.status_code == 201, preview.data
    assert preview.data["rows"][0]["errors"] == {}
    assert preview.data["rows"][0]["data"]["barcode"] == "0012345678905"
