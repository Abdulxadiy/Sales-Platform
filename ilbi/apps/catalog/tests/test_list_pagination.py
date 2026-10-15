"""List pages must fetch bounded rows without losing filters or tenant scope."""
import io
from datetime import timedelta

import pytest
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.inventory.models import Stock, StockBatch, StockMovement
from apps.permissions.models import Permission
from tests.factories import CategoryFactory, EmployeeFactory, ProductFactory, ProductVariantFactory, StaffFactory

pytestmark = pytest.mark.django_db
VARIANTS = '/api/v1/catalog/variants/'
STOCK = '/api/v1/inventory/stock/'
MOVEMENTS = '/api/v1/inventory/movements/'


@pytest.fixture
def client(owner, tenant):
    client = APIClient()
    client.force_authenticate(owner)
    return client


@pytest.fixture
def rows(tenant, owner):
    category = CategoryFactory(tenant=tenant)
    variants = []
    for index in range(31):
        product = ProductFactory(tenant=tenant, category=category, name=f'Item {index:03}')
        variant = ProductVariantFactory(tenant=tenant, product=product, code=f'CODE-{index}', barcode=f'BAR-{index}')
        variants.append(variant)
        Stock.objects.create(tenant=tenant, product_variant=variant, quantity=index + 1, last_cost_price=index + 1)
        StockMovement.objects.create(tenant=tenant, product_variant=variant, created_by=owner,
                                     type='kirim', direction='in', quantity=index + 1)
    # A different tenant must never be exposed by lookup or ids.
    variants.append(ProductVariantFactory())
    return variants


def test_variant_page_limits_related_queries_and_preserves_expiry(client, rows, tenant):
    StockBatch.objects.create(tenant=tenant, product_variant=rows[0], quantity=2,
                              expires_on=timezone.localdate() - timedelta(days=1))
    with CaptureQueriesContext(connection) as captured:
        response = client.get(VARIANTS, {'page': 1, 'page_size': 20})
    assert response.status_code == 200
    assert response.data['count'] == 31
    assert len(response.data['results']) == 20
    assert response.data['results'][0]['expired_batches'][0]['quantity'] == '2.000'
    assert len(captured) <= 9  # bounded joins/prefetches, never one query per row
    assert any('LIMIT 20' in q['sql'] for q in captured)
    second = client.get(VARIANTS, {'page': 2, 'page_size': 20}).data
    assert len(second['results']) == 11
    assert not ({v['id'] for v in response.data['results']} & {v['id'] for v in second['results']})
    filtered = client.get(VARIANTS, {'page': 1, 'search': 'Item 030'}).data
    assert filtered['count'] == 1
    assert filtered['results'][0]['id'] == rows[30].pk


@pytest.mark.parametrize('lookup', ['bar-30', 'code-30'])
def test_scanner_finds_items_outside_first_page(client, rows, lookup):
    result = client.get(VARIANTS, {'page': 1, 'lookup': lookup}).data
    assert result['count'] == 1
    assert result['results'][0]['id'] == rows[30].pk


def test_checkout_only_fetches_requested_active_variants_in_tenant(client, rows):
    result = client.get(VARIANTS, {'ids': f'{rows[1].pk},{rows[30].pk},{rows[-1].pk}'}).data
    assert {v['id'] for v in result} == {rows[1].pk, rows[30].pk}
    rows[1].is_active = False
    rows[1].save(update_fields=['is_active'])
    result = client.get(VARIANTS, {'ids': f'{rows[1].pk},{rows[30].pk}'}).data
    assert [v['id'] for v in result] == [rows[30].pk]
    assert client.get(VARIANTS, {'ids': '1,invalid'}).status_code == 400
    assert client.get(VARIANTS, {'lookup': rows[-1].sku, 'page': 1}).data['count'] == 0


def test_stock_filters_sort_before_pagination_and_export_all_matches(client, rows):
    params = {'page': 2, 'page_size': 20, 'sort': 'qty_desc', 'category': rows[0].product.category_id}
    result = client.get(STOCK, params).data
    assert result['count'] == 31
    assert len(result['results']) == 11
    assert result['results'][0]['quantity'] == '11.000'
    assert result['results'][-1]['quantity'] == '1.000'
    result = client.get(STOCK, {'page': 1, 'search': 'BAR-30'}).data
    assert result['count'] == 1
    export = client.get('/api/v1/inventory/stock/export/', params)
    assert export.status_code == 200
    sheet = load_workbook(io.BytesIO(export.content), read_only=True).active
    exported = list(sheet.values)
    assert len(exported) == 32  # all matching rows, not just page two
    assert [row[6] for row in exported[1:]] == list(range(31, 0, -1))
    empty_category = CategoryFactory(tenant=rows[0].tenant)
    assert client.get(STOCK, {'category': empty_category.pk, 'page': 1}).data['count'] == 0
    assert client.get(STOCK, {'category': 'bad', 'page': 1}).status_code == 400


def test_full_stock_export_ignores_filters_and_pages_but_keeps_tenant_scope(client, rows):
    Stock.objects.filter(product_variant=rows[0]).update(quantity=0)
    rows[0].is_active = False
    rows[0].save(update_fields=['is_active'])
    Stock.objects.create(tenant=rows[-1].tenant, product_variant=rows[-1], quantity=99)
    response = client.get('/api/v1/inventory/stock/export/', {
        'scope': 'all', 'page': 2, 'page_size': 1, 'search': 'BAR-30',
        'category': rows[-1].product.category_id,
    })
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    exported = list(workbook.active.values)
    assert len(exported) == 32
    headers = exported[0]
    assert {row[3] for row in exported[1:]} == {v.sku for v in rows[:-1]}
    assert exported[1][6] == 0
    assert exported[1][headers.index('Holat')] == 'Arxivlangan'
    metadata = dict(workbook['Hisobot'].values)
    assert metadata['Qamrov'] == 'Butun ombor'
    assert metadata['Qatorlar soni'] == 31
    assert 'Qidiruv' not in metadata
    assert client.get('/api/v1/inventory/stock/export/', {'scope': 'invalid'}).status_code == 400


def test_movements_filters_and_pagination(client, rows):
    StockMovement.objects.filter(product_variant=rows[0]).update(created_at=timezone.now() - timedelta(days=3))
    first = client.get(MOVEMENTS, {'page': 1, 'page_size': 20}).data
    second = client.get(MOVEMENTS, {'page': 2, 'page_size': 20}).data
    assert first['count'] == 31
    assert len(first['results']) == 20
    assert len(second['results']) == 11
    assert not ({v['id'] for v in first['results']} & {v['id'] for v in second['results']})
    result = client.get(MOVEMENTS, {'page': 1, 'search': 'Item 030', 'type': 'kirim', 'date': 'today'}).data
    assert result['count'] == 1
    assert client.get(MOVEMENTS, {'page': 1, 'date': 'today'}).data['count'] == 30
    assert client.get(MOVEMENTS, {'page': 1, 'type': 'sotuv'}).data['count'] == 0
    assert isinstance(client.get(MOVEMENTS).data, list)  # legacy clients still supported


def test_staff_cost_permission_checked_once_per_request(client, rows, tenant):
    user = StaffFactory()
    employee = EmployeeFactory(user=user, tenant=tenant)
    employee.permissions.add(Permission.objects.get(category='inventory', codename='view_stock'))
    client.force_authenticate(user)
    for endpoint in (STOCK, MOVEMENTS):
        with CaptureQueriesContext(connection) as small:
            response = client.get(endpoint, {'page': 1, 'page_size': 1})
        assert response.status_code == 200
        with CaptureQueriesContext(connection) as large:
            response = client.get(endpoint, {'page': 1, 'page_size': 20})
        assert response.status_code == 200
        assert len(large) == len(small)
        assert all('last_cost_price' not in row and 'cost_price' not in row for row in response.data['results'])
    employee.permissions.add(Permission.objects.get(category='catalog', codename='view_cost_price'))
    response = client.get(STOCK, {'page': 1})
    assert 'last_cost_price' in response.data['results'][0]  # no authorization cache across requests
