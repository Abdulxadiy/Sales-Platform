import io
from datetime import timedelta

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.catalog.models import Product, ProductVariant
from apps.catalog.models.import_model import CatalogImportJob
from apps.inventory.models import StockBatch
from apps.permissions.models import Permission
from tests.factories import (
    CategoryFactory,
    EmployeeFactory,
    ProductVariantFactory,
    SaleFactory,
    StaffFactory,
    TenantFactory,
)

pytestmark = pytest.mark.django_db
PREVIEW = "/api/v1/catalog/import/preview/"


@pytest.fixture
def client(owner, tenant):
    client = APIClient()
    client.force_authenticate(owner)
    return client


def csv_file(body):
    return SimpleUploadedFile("catalog.csv", body.encode("utf-8-sig"), content_type="text/csv")


def excel_file(rows, barcode_format=None):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    if barcode_format:
        workbook.active["D2"].number_format = barcode_format
    output = io.BytesIO()
    workbook.save(output)
    return SimpleUploadedFile("catalog.xlsx", output.getvalue())


def contents(response):
    assert response.status_code == 200, getattr(response, "data", "")
    return list(load_workbook(io.BytesIO(response.content), data_only=False).active.values)


def test_csv_preview_does_not_change_catalog_and_commit_is_idempotent(client, tenant):
    expiry = (timezone.localdate() + timedelta(days=30)).isoformat()
    response = client.post(
        PREVIEW,
        {
            "file": csv_file(
                "mahsulot,kategoriya,sotuv_narxi,miqdor,tannarx,partiya,yaroqlilik_sanasi,shtrix_kod\n"
                "Sut,Sut mahsulotlari,15000,4,9000,LOT-1," + expiry + ",001234\n"
            )
        },
        format="multipart",
    )
    assert response.status_code == 201, response.data
    assert response.data["rows"][0]["errors"] == {}
    assert not Product.objects.filter(tenant=tenant).exists()
    url = f"/api/v1/catalog/import/{response.data['id']}/commit/"
    result = client.post(url)
    assert result.status_code == 200, result.data
    assert result.data["created_count"] == 1
    batch = StockBatch.objects.get(tenant=tenant)
    assert batch.batch_code == "LOT-1"
    assert batch.expires_on.isoformat() == expiry
    assert batch.quantity == 4
    assert batch.product_variant.barcode == "001234"
    assert client.post(url).data == result.data
    assert StockBatch.objects.filter(tenant=tenant).count() == 1


def test_duplicate_and_invalid_rows_block_entire_import(client, tenant):
    response = client.post(
        PREVIEW,
        {
            "file": csv_file(
                "mahsulot,kategoriya,sotuv_narxi,shtrix_kod\nA,Test,5000,AB1\nB,Test,-1,AB2\nC,Test,6000,AB1\n"
            )
        },
    )
    assert response.status_code == 201
    assert "price_recommended" in response.data["rows"][1]["errors"]
    assert "barcode" in response.data["rows"][2]["errors"]
    result = client.post(f"/api/v1/catalog/import/{response.data['id']}/commit/")
    assert result.status_code == 400
    assert not Product.objects.filter(tenant=tenant).exists()


def test_import_adds_multiple_variants_to_one_product_and_flags_archived_category(client, tenant):
    data = "mahsulot,kategoriya,sotuv_narxi,variant\nKo‘ylak,Kiyim,10000,Oq\nKo‘ylak,Kiyim,12000,Qora\n"
    preview = client.post(PREVIEW, {"file": csv_file(data)})
    assert not any(row["errors"] for row in preview.data["rows"])
    commit = client.post(f"/api/v1/catalog/import/{preview.data['id']}/commit/")
    assert commit.status_code == 200, commit.data
    assert commit.data["created_count"] == 2
    assert Product.objects.filter(tenant=tenant).count() == 1
    assert ProductVariant.objects.filter(tenant=tenant).count() == 2
    CategoryFactory(tenant=tenant, name="Arxiv", is_active=False)
    preview = client.post(PREVIEW, {"file": csv_file("mahsulot,kategoriya,sotuv_narxi\nA,Arxiv,100\n")})
    assert "category" in preview.data["rows"][0]["errors"]


def test_commit_rechecks_catalog_and_job_owner(client, tenant):
    response = client.post(PREVIEW, {"file": csv_file("mahsulot,kategoriya,sotuv_narxi,shtrix_kod\nA,Test,5000,AB1\n")})
    url = f"/api/v1/catalog/import/{response.data['id']}/commit/"
    ProductVariantFactory(tenant=tenant, barcode="AB1")
    assert client.post(url).status_code == 400
    staff = StaffFactory()
    employee = EmployeeFactory(user=staff, tenant=tenant)
    employee.permissions.add(Permission.objects.get(category="catalog", codename="add_product"))
    client.force_authenticate(staff)
    assert client.post(url).status_code == 404


def test_import_permissions_are_checked_per_row_and_again_at_commit(tenant):
    category = CategoryFactory(tenant=tenant)
    user = StaffFactory()
    employee = EmployeeFactory(user=user, tenant=tenant)
    permission = Permission.objects.get(category="catalog", codename="add_product")
    employee.permissions.add(permission)
    client = APIClient()
    client.force_authenticate(user)
    body = f"mahsulot,kategoriya,sotuv_narxi,miqdor,tannarx\nA,{category.name},100,2,10\nB,Nova,200,0,0\n"
    preview = client.post(PREVIEW, {"file": csv_file(body)})
    assert preview.status_code == 201, preview.data
    assert "quantity" in preview.data["rows"][0]["errors"]
    assert "category" in preview.data["rows"][1]["errors"]
    body = f"mahsulot,kategoriya,sotuv_narxi\nA,{category.name},100\n"
    preview = client.post(PREVIEW, {"file": csv_file(body)})
    assert preview.data["rows"][0]["errors"] == {}
    employee.permissions.clear()
    assert client.post(f"/api/v1/catalog/import/{preview.data['id']}/commit/").status_code == 403
    assert not Product.objects.filter(tenant=tenant).exists()


def test_excel_import_preserves_zero_formatted_barcode_and_rejects_formulas(client):
    upload = excel_file(
        [["mahsulot", "kategoriya", "sotuv_narxi", "shtrix_kod"], ["A", "Test", 100, 1234]], barcode_format="000000"
    )
    response = client.post(PREVIEW, {"file": upload})
    assert response.status_code == 201
    assert response.data["rows"][0]["data"]["barcode"] == "001234"
    spaced = excel_file([["mahsulot", "kategoriya", "sotuv_narxi"], [], ["A", "Test", 100]])
    assert client.post(PREVIEW, {"file": spaced}).data['rows'][0]['row_number'] == 3
    formula = excel_file([["mahsulot", "kategoriya", "sotuv_narxi"], ["A", "Test", "=100+10"]])
    assert client.post(PREVIEW, {"file": formula}).status_code == 400


def test_bad_format_large_import_and_expired_preview_are_rejected(client):
    assert client.post(PREVIEW, {"file": SimpleUploadedFile("fake.xlsx", b"not excel")}).status_code == 400
    body = "mahsulot,kategoriya,sotuv_narxi\n" + "A,Test,100\n" * 1001
    assert client.post(PREVIEW, {"file": csv_file(body)}).status_code == 400
    upload = excel_file([["mahsulot", "kategoriya", "sotuv_narxi"], *[["A", "Test", 100] for _ in range(1001)]])
    assert client.post(PREVIEW, {"file": upload}).status_code == 400
    preview = client.post(PREVIEW, {"file": csv_file("mahsulot,kategoriya,sotuv_narxi\nA,Test,100\n")})
    CatalogImportJob.objects.filter(pk=preview.data["id"]).update(expires_at=timezone.now() - timedelta(seconds=1))
    assert client.post(f"/api/v1/catalog/import/{preview.data['id']}/commit/").status_code == 400


def test_excel_export_respects_selection_tenant_and_owner_only_access(client, owner, tenant):
    from apps.inventory.services import StockService

    variant = ProductVariantFactory(tenant=tenant)
    StockService.intake(tenant=tenant, product_variant=variant, quantity=3, cost_price=100, created_by=owner)
    stock = variant.stock
    assert "Tannarx" in contents(client.post("/api/v1/inventory/stock/export/", {"ids": [stock.pk]}, format="json"))[0]
    staff = StaffFactory()
    employee = EmployeeFactory(user=staff, tenant=tenant)
    employee.permissions.add(Permission.objects.get(category="inventory", codename="view_stock"))
    client.force_authenticate(staff)
    assert client.post("/api/v1/inventory/stock/export/", {"ids": [stock.pk]}, format="json").status_code == 403
    assert client.post("/api/v1/inventory/stock/export/", {"ids": []}, format="json").status_code == 403
    batches = client.get("/api/v1/inventory/batches/")
    assert batches.status_code == 200
    assert "cost_price" not in batches.data[0]
    assert client.get("/api/v1/sales/export/").status_code == 403
    client.force_authenticate(owner)
    assert len(contents(client.post("/api/v1/inventory/stock/export/", {"ids": []}, format="json"))) == 1
    assert client.post("/api/v1/inventory/stock/export/", {"ids": [99999999]}, format="json").status_code == 400


def test_sales_export_includes_all_filtered_pages_and_escapes_excel_formula(client, owner, tenant):
    other = TenantFactory()
    SaleFactory(tenant=other, receipt_number="FILTER-FOREIGN")
    SaleFactory.create_batch(28, tenant=tenant, sold_by=owner)
    selected = SaleFactory(tenant=tenant, sold_by=owner, receipt_number="FILTER-LOCAL")
    rows = contents(client.get("/api/v1/sales/export/?page=2"))
    assert len(rows) == 30
    rows = contents(client.get("/api/v1/sales/export/?search=FILTER"))
    assert len(rows) == 2
    assert rows[1][0] == selected.receipt_number
    from apps.sales.models import Counterparty

    cp = Counterparty.objects.create(tenant=tenant, name='=HYPERLINK("bad")')
    response = client.post("/api/v1/sales/counterparties/export/", {"ids": [cp.pk]}, format="json")
    cell = load_workbook(io.BytesIO(response.content)).active["A2"]
    assert cell.data_type == "s"
    assert cell.value == cp.name


def test_generated_barcode_is_idempotent_and_tenant_scoped(client, tenant):
    variant = ProductVariantFactory(tenant=tenant)
    url = f"/api/v1/catalog/variants/{variant.pk}/barcode/"
    result = client.post(url)
    assert result.status_code == 200, result.data
    assert result.data["barcode"].isdigit()
    assert len(result.data["barcode"]) == 13
    digits = result.data["barcode"]
    assert sum(int(digit) * (1 if i % 2 == 0 else 3) for i, digit in enumerate(digits)) % 10 == 0
    assert client.post(url).data["barcode"] == result.data["barcode"]
    from apps.core.models import AuditLog

    assert AuditLog.objects.filter(tenant=tenant, action="barcode_generate", target_id=str(variant.pk)).count() == 1
    other = ProductVariantFactory()
    assert client.post(f"/api/v1/catalog/variants/{other.pk}/barcode/").status_code == 404


@pytest.mark.parametrize("existing_code", ["0012345678905", "INV-5-20002", "INV-5-20002-1"])
def test_existing_barcode_is_not_replaced(client, tenant, existing_code):
    variant = ProductVariantFactory(tenant=tenant, barcode=existing_code)
    result = client.post(f"/api/v1/catalog/variants/{variant.pk}/barcode/")
    assert result.status_code == 200, result.data
    assert result.data["barcode"] == existing_code
    variant.refresh_from_db()
    assert variant.barcode == existing_code


def test_numeric_barcode_generator_skips_collisions(client, tenant, monkeypatch):
    ProductVariantFactory(tenant=tenant, barcode="2000000000015")
    candidates = iter([1, 2])
    monkeypatch.setattr("api.v1.catalog.views.label_views.secrets.randbelow", lambda _: next(candidates))
    variant = ProductVariantFactory(tenant=tenant, barcode=None)
    result = client.post(f"/api/v1/catalog/variants/{variant.pk}/barcode/")
    assert result.status_code == 200, result.data
    assert result.data["barcode"] == "2000000000022"
