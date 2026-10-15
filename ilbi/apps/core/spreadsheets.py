import csv
import io
import math
import re
import textwrap
from datetime import date, datetime
from decimal import Decimal
from itertools import chain, islice
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet
from rest_framework.exceptions import ValidationError

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_IMPORT_ROWS = 1000
MAX_EXPORT_ROWS = 20000
MAX_EXCEL_DATA_ROWS = 1048575  # One worksheet, excluding its header.
CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
MONEY_FORMAT = "#,##0.00;[Red](#,##0.00);0.00"
QUANTITY_FORMAT = "#,##0.000;[Red](#,##0.000);0"
INTEGER_FORMAT = "#,##0;[Red](#,##0);0"
NAVY = "17324D"
TEAL = "007F86"
CENTERED_COLUMNS = {
    "SKU",
    "Shtrix-kod",
    "Mahsulot kodi",
    "Birlik",
    "Qoldiq",
    "Valyuta",
    "Holat",
    "Chek",
    "Telefon",
    "Vaqt (Toshkent)",
    "shtrix_kod",
    "birlik",
    "valyuta",
    "miqdor",
    "partiya",
    "yaroqlilik_sanasi",
}


def _quantity_format(value):
    # Optional decimal placeholders leave a dangling separator in Excel/LibreOffice.
    # Choose a concrete format so 181, 0 and 1.125 all stay visible and unambiguous.
    places = max(0, min(3, -Decimal(str(value)).normalize().as_tuple().exponent))
    if not places:
        return INTEGER_FORMAT
    decimal_part = "." + "0" * places
    return f"#,##0{decimal_part};[Red](#,##0{decimal_part});0"


def _cell(sheet, value, number_format=None):
    if isinstance(value, datetime) and timezone.is_aware(value):
        value = timezone.localtime(value).replace(tzinfo=None)
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.value = CONTROL_CHARS.sub("", value)
        cell.data_type = "s"  # User-supplied formula-like strings stay literal, including on the summary.
    if number_format == QUANTITY_FORMAT and isinstance(value, (int, float, Decimal)):
        number_format = _quantity_format(value)
    cell.number_format = number_format or (
        "dd.mm.yyyy hh:mm"
        if isinstance(value, datetime)
        else "dd.mm.yyyy"
        if isinstance(value, date)
        else MONEY_FORMAT
        if isinstance(value, (Decimal, float))
        else "General"
    )
    return cell


def _append_sized_row(sheet, cells, widths, row_number, *, minimum=32):
    """Set row height before streaming it, then discard its dimension to bound memory."""
    lines = 1
    for cell, width in zip(cells, widths, strict=True):
        if cell.value is None or not cell.alignment.wrap_text:
            continue
        # Reserve room for cell padding and use a conservative estimate for proportional fonts.
        capacity = max(1, math.floor((width - 4) * 0.9))
        wrapped = sum(max(1, len(textwrap.wrap(part, capacity))) for part in str(cell.value).split("\n"))
        lines = max(lines, wrapped)
    sheet.row_dimensions[row_number].height = min(409, max(minimum, lines * 16 + 14))
    sheet.append(cells)
    del sheet.row_dimensions[row_number]


def _column_widths(headers, sample, formats, overrides):
    widths = []
    for column, header in enumerate(headers):
        if header in overrides:
            widths.append(overrides[header])
            continue
        longest = len(str(header))
        for row in sample:
            value = row[column]
            if isinstance(value, (int, float, Decimal)):
                text = format(value, ",.3f" if formats.get(header) == QUANTITY_FORMAT else ",.2f")
            elif isinstance(value, datetime):
                text = "31.12.2026 23:59"
            else:
                text = str(value or "")
            longest = max(longest, *(len(part) for part in text.split("\n")))
        # A bounded sample sets readable widths without consuming the entire export.
        widths.append(min(64, max(18, longest + 6)))
    return widths


def _page_layout(sheet, title, landscape=False):
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = Worksheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.oddHeader.left.text = "ILBI | " + title
    sheet.oddHeader.left.size = 10
    sheet.oddFooter.left.text = "ILBI · " + timezone.localtime().strftime("%d.%m.%Y %H:%M")
    sheet.oddFooter.right.text = "&P / &N"


def _report_sheet(workbook, title, report, row_count, totals):
    sheet = workbook.create_sheet(report.get("sheet_title", "Hisobot"))
    _page_layout(sheet, title)
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 65
    sheet.sheet_format.defaultRowHeight = 42
    sheet.row_dimensions[1].height = 48
    sheet.freeze_panes = "A4"
    sheet.print_options.horizontalCentered = False
    header_fill = PatternFill("solid", fgColor=NAVY)
    row_number = 0

    def line(label, value=None, *, heading=False, number_format=None):
        nonlocal row_number
        row_number += 1
        cells = [_cell(sheet, label), _cell(sheet, value, number_format)]
        for cell in cells:
            cell.font = Font(name="Calibri", size=11, color="FFFFFF" if heading else NAVY, bold=heading)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
            if heading:
                cell.fill = header_fill
        if cells[1].data_type == "n":
            cells[1].alignment = Alignment(horizontal="right", vertical="center", indent=1, shrink_to_fit=True)
        _append_sized_row(sheet, cells, [38, 65], row_number, minimum=48 if heading else 36)

    line("ILBI", title, heading=True)
    line("Do‘kon", report.get("tenant", ""))
    line("Tayyorlangan vaqt (Toshkent)", timezone.localtime())
    line("Qamrov", report.get("scope", "Barcha qatorlar"))
    line("Qatorlar soni", row_count, number_format="#,##0")
    for label, value in report.get("details", []):
        line(label, value)
    for group, values in zip(report.get("groups", []), totals, strict=True):
        line("")
        line(group["title"], "Jami", heading=True)
        for key, sums in sorted(values.items(), key=lambda item: str(item[0])):
            for column, label in group["columns"].items():
                line(f"{label} · {key}", sums[column], number_format=group.get("format", MONEY_FORMAT))
        if not values:
            line("Ma’lumot yo‘q", "Tanlangan qamrovda qatorlar mavjud emas.")
    if row_count == 0 and not report.get("notes"):
        line("Ma’lumot yo‘q", "Tanlangan qamrovda qatorlar mavjud emas.")
    if report.get("notes"):
        line("")
        line("Qo‘llanma", "Tavsif", heading=True)
        for label, description in report["notes"]:
            line(label, description)


def excel_response(headers, rows, filename, title="ILBI", *, formats=None, widths=None, report=None):
    """Stream a styled data sheet and an optional summary without buffering all rows."""
    workbook = Workbook(write_only=True)
    workbook.properties.creator = "ILBI"
    workbook.properties.title = title
    workbook.properties.subject = report.get("tenant", "") if report else title
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = "B2"
    sheet.sheet_properties.tabColor = TEAL
    sheet.sheet_format.defaultRowHeight = 24
    sheet.row_dimensions[1].height = 36
    _page_layout(sheet, title, landscape=len(headers) > 6)
    sheet.print_title_rows = "1:1"
    formats = formats or {}
    rows = iter(rows)
    sample = list(islice(rows, 200))
    column_widths = _column_widths(headers, sample, formats, widths or {})
    for column, header in enumerate(headers, start=1):
        dimension = sheet.column_dimensions[get_column_letter(column)]
        dimension.width = column_widths[column - 1]
        if header in formats:
            dimension.number_format = formats[header]
    totals = [{} for _ in (report or {}).get("groups", [])]
    header_fill = PatternFill("solid", fgColor=NAVY)
    stripe_fill = PatternFill("solid", fgColor="F0F5FA")
    edge = Side(style="thin", color="CCD8E4")
    border = Border(left=edge, right=edge, bottom=edge)
    body_font = Font(name="Calibri", size=11, color=NAVY)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    name_font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    text_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    number_alignment = Alignment(horizontal="right", vertical="center", indent=1, shrink_to_fit=True)
    centered_number_alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    alignments = []
    for column, header in enumerate(headers):
        numeric = formats.get(header) in (MONEY_FORMAT, QUANTITY_FORMAT) or any(
            isinstance(row[column], (int, float, Decimal)) for row in sample
        )
        alignments.append(
            center_alignment if header in CENTERED_COLUMNS else number_alignment if numeric else text_alignment
        )
    row_count = 0
    for index, values in enumerate(chain([headers], sample, rows)):
        cells = []
        for column, value in enumerate(values):
            cell = _cell(sheet, value, formats.get(headers[column]) if index else None)
            cell.font = name_font if column == 0 else body_font
            cell.alignment = alignments[column]
            if index and cell.data_type == "n" and cell.alignment.horizontal == "center":
                cell.alignment = centered_number_alignment
            cell.border = border
            if index == 0:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(
                    horizontal=alignments[column].horizontal,
                    vertical="center",
                    wrap_text=True,
                    indent=2 if alignments[column].horizontal == "right" else alignments[column].indent,
                )
            elif index % 2 == 0:
                cell.fill = stripe_fill
            cells.append(cell)
        _append_sized_row(sheet, cells, column_widths, index + 1, minimum=44 if index == 0 else 32)
        if index:
            row_count += 1
            for group, grouped_totals in zip((report or {}).get("groups", []), totals, strict=True):
                key = values[group["key"]] if group.get("key") is not None else group.get("label", "Jami")
                sums = grouped_totals.setdefault(key, {column: Decimal("0") for column in group["columns"]})
                for column in sums:
                    if values[column] is not None:
                        sums[column] += Decimal(str(values[column]))
    last_column = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A1:{last_column}{row_count + 1}"
    sheet.print_area = f"A1:{last_column}{row_count + 1}"
    if report:
        _report_sheet(workbook, title, report, row_count, totals)
    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type=XLSX_MIME)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["Cache-Control"] = "private, no-store"
    return response


def read_import_file(upload):
    if upload.size > 8 * 1024 * 1024:
        raise ValidationError("Fayl 8 MB dan kichik bo‘lishi kerak.")
    suffix = Path(upload.name).suffix.lower()
    content = upload.read()
    rows = []
    try:
        if suffix == ".csv":
            text = content.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(text), dialect)
            for row in reader:
                rows.append(row)
                if len(rows) > MAX_IMPORT_ROWS + 1:
                    raise ValidationError("Bitta faylda ko‘pi bilan 1000 qator bo‘lsin.")
        elif suffix == ".xlsx":
            with ZipFile(io.BytesIO(content)) as archive:
                if (
                    len(archive.infolist()) > 1000
                    or sum(entry.file_size for entry in archive.infolist()) > 32 * 1024 * 1024
                ):
                    raise ValidationError("Excel faylining ochilgan hajmi juda katta.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.active
                if sheet is None:
                    raise ValidationError("Excel faylida jadval yo‘q.")
                if sheet.max_column and sheet.max_column > 20:
                    raise ValidationError("Faylda ko‘pi bilan 20 ustun bo‘lsin.")
                if sheet.max_row and sheet.max_row > MAX_IMPORT_ROWS + 1:
                    raise ValidationError(
                        "Bitta faylda ko‘pi bilan 1000 qator bo‘lsin; ortiqcha bo‘sh qatorlarni olib tashlang."
                    )
                sheet.reset_dimensions()
                for number, row in enumerate(sheet.iter_rows(max_col=21), start=1):
                    if number > MAX_IMPORT_ROWS + 1:
                        raise ValidationError("Bitta faylda ko‘pi bilan 1000 qator bo‘lsin.")
                    values = []
                    for cell in row:
                        if cell.data_type == "f":
                            raise ValidationError("Formulalar o‘rniga qiymatlarni saqlang: Paste special → Values.")
                        value = cell.value
                        if isinstance(value, (date, datetime)):
                            value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                        elif isinstance(value, (int, float)):
                            if isinstance(value, float) and value.is_integer():
                                value = int(value)
                            # Preserve explicitly formatted leading zeros in barcode cells.
                            if str(cell.number_format).strip("0") == "" and cell.number_format:
                                value = str(value).zfill(len(cell.number_format))
                        values.append("" if value is None else str(value))
                    while values and values[-1] == "":
                        values.pop()
                    rows.append(values)
                    if len(rows) > MAX_IMPORT_ROWS + 1:
                        raise ValidationError("Bitta faylda ko‘pi bilan 1000 qator bo‘lsin.")
            finally:
                workbook.close()
        else:
            raise ValidationError("Faqat .csv yoki .xlsx faylni tanlang.")
    except (
        UnicodeError,
        csv.Error,
        BadZipFile,
        KeyError,
        ValueError,
        OSError,
        TypeError,
        IndexError,
        ParseError,
        DefusedXmlException,
        InvalidFileException,
    ) as exc:
        raise ValidationError("Fayl o‘qilmadi. UTF-8 CSV yoki haqiqiy .xlsx faylni tanlang.") from exc
    if not rows or not any(any(row) for row in rows):
        raise ValidationError("Fayl bo‘sh.")
    if len(rows[0]) > 20:
        raise ValidationError("Faylda ko‘pi bilan 20 ustun bo‘lsin.")
    return rows
