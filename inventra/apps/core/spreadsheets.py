import csv
import io
import re
from datetime import date, datetime
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework.exceptions import ValidationError

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_IMPORT_ROWS = 1000
MAX_EXPORT_ROWS = 20000
CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def excel_response(headers, rows, filename, title="Inventra"):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = "A2"
    for column, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = min(40, max(16, len(str(header)) + 3))
    for index, values in enumerate([headers, *rows]):
        cells = []
        for value in values:
            if isinstance(value, datetime) and timezone.is_aware(value):
                value = timezone.localtime(value).replace(tzinfo=None)
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.value = CONTROL_CHARS.sub("", value)
                cell.data_type = "s"  # Names such as '=HYPERLINK(...)' remain literal text.
            if isinstance(value, datetime):
                cell.number_format = 'yyyy-mm-dd hh:mm'
            if index == 0:
                cell.font = Font(bold=True)
            cells.append(cell)
        sheet.append(cells)
    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type=XLSX_MIME)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["Cache-Control"] = "private, no-store"
    return response


def read_import_file(upload):
    if upload.size > 8 * 1024 * 1024:
        raise ValidationError("Fayl 8 MB dan kichik bo‘lishi kerak.")
    suffix = Path(upload.name).suffix.lower()
    content = upload.read()
    rows = []
    try:
        if suffix == ".csv":
            text = content.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(text), dialect)
            for row in reader:
                rows.append(row)
                if len(rows) > MAX_IMPORT_ROWS + 1:
                    raise ValidationError("Bitta faylda ko‘pi bilan 1000 qator bo‘lsin.")
        elif suffix == ".xlsx":
            with ZipFile(io.BytesIO(content)) as archive:
                if (
                    len(archive.infolist()) > 1000
                    or sum(entry.file_size for entry in archive.infolist()) > 32 * 1024 * 1024
                ):
                    raise ValidationError("Excel faylining ochilgan hajmi juda katta.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.active
                if sheet is None:
                    raise ValidationError("Excel faylida jadval yo‘q.")
                if sheet.max_column and sheet.max_column > 20:
                    raise ValidationError("Faylda ko‘pi bilan 20 ustun bo‘lsin.")
                if sheet.max_row and sheet.max_row > MAX_IMPORT_ROWS + 1:
                    raise ValidationError(
                        "Bitta faylda ko‘pi bilan 1000 qator bo‘lsin; ortiqcha bo‘sh qatorlarni olib tashlang."
                    )
                sheet.reset_dimensions()
                for number, row in enumerate(sheet.iter_rows(max_col=21), start=1):
                    if number > MAX_IMPORT_ROWS + 1:
                        raise ValidationError("Bitta faylda ko‘pi bilan 1000 qator bo‘lsin.")
                    values = []
                    for cell in row:
                        if cell.data_type == "f":
                            raise ValidationError("Formulalar o‘rniga qiymatlarni saqlang: Paste special → Values.")
                        value = cell.value
                        if isinstance(value, (date, datetime)):
                            value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                        elif isinstance(value, (int, float)):
                            if isinstance(value, float) and value.is_integer():
                                value = int(value)
                            # Preserve explicitly formatted leading zeros in barcode cells.
                            if str(cell.number_format).strip("0") == "" and cell.number_format:
                                value = str(value).zfill(len(cell.number_format))
                        values.append("" if value is None else str(value))
                    while values and values[-1] == "":
                        values.pop()
                    rows.append(values)
                    if len(rows) > MAX_IMPORT_ROWS + 1:
                        raise ValidationError("Bitta faylda ko‘pi bilan 1000 qator bo‘lsin.")
            finally:
                workbook.close()
        else:
            raise ValidationError("Faqat .csv yoki .xlsx faylni tanlang.")
    except (
        UnicodeError,
        csv.Error,
        BadZipFile,
        KeyError,
        ValueError,
        OSError,
        TypeError,
        IndexError,
        ParseError,
        DefusedXmlException,
        InvalidFileException,
    ) as exc:
        raise ValidationError("Fayl o‘qilmadi. UTF-8 CSV yoki haqiqiy .xlsx faylni tanlang.") from exc
    if not rows or not any(any(row) for row in rows):
        raise ValidationError("Fayl bo‘sh.")
    if len(rows[0]) > 20:
        raise ValidationError("Faylda ko‘pi bilan 20 ustun bo‘lsin.")
    return rows
